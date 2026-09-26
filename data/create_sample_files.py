"""
Generates sample assignment files (.txt and .docx) in data/samples/
for immediate testing in the Streamlit UI.
"""

import os
from docx import Document

samples_dir = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(samples_dir, exist_ok=True)

# 1. AI Original
text_ai_original = """Artificial Intelligence and Machine Learning Fundamentals
Student Name: Rahul Sharma
Roll Number: CS-2024-042

1. Introduction
Machine learning is a subset of artificial intelligence that enables computers to learn from historical data and make automated predictions without explicit rule programming. It can be divided into supervised learning, unsupervised learning, and reinforcement learning.

2. Supervised Learning
Supervised learning algorithms are trained using labeled datasets where the desired ground truth is known. Examples include linear regression, decision trees, support vector machines, and deep neural networks.

3. Neural Networks and Deep Learning
Convolutional Neural Networks are designed specifically for image recognition and computer vision tasks. They extract hierarchical spatial patterns through convolutional layers, pooling layers, and fully connected activation layers.

4. Conclusion
Machine learning has transformed automated decision making across healthcare diagnostics, autonomous vehicles, and natural language processing systems."""

# 2. AI Exact Copy
text_ai_duplicate = text_ai_original

# 3. AI Paraphrased (Similar meaning, modified wording)
text_ai_paraphrased = """Overview of Machine Learning and Artificial Intelligence
Student Name: Amit Verma
Roll Number: CS-2024-058

1. Overview
Machine learning represents a specialized branch within AI that allows computational systems to extract patterns from training data to boost precision without hard-coded logic. Its primary paradigms comprise supervised learning, unsupervised learning, and reinforcement learning.

2. Supervised Learning Techniques
Supervised approaches utilize annotated data sets with known ground truth targets. Popular computational models encompass linear regression, decision trees, support vector machines, and artificial neural networks.

3. Deep Learning Architectures
Convolutional neural network architectures are tailored for digital image recognition and visual analysis. They capture structural spatial hierarchies using alternating convolutional filtering, pooling reductions, and dense layers.

4. Summary
Modern machine learning models continue to revolutionize automated problem solving across medical diagnostics, self-driving transportation, and human language comprehension."""

# 4. DBMS (Completely Different)
text_dbms = """Database Management Systems and Normalization
Student Name: Priya Patel
Roll Number: CS-2024-019

1. Relational Database Concepts
A relational database management system organizes data into tables consisting of rows and columns with primary keys and foreign keys enforcing referential integrity across relational schemas.

2. Transaction Reliability and ACID
ACID properties ensure database transaction reliability: Atomicity guarantees all-or-nothing execution, Consistency enforces semantic integrity, Isolation prevents concurrent collision, and Durability ensures committed updates persist permanently.

3. Normalization Techniques
Database normalization systematically restructures relations to eliminate redundant storage and insertion, update, or deletion anomalies across First, Second, Third, and Boyce-Codd Normal Forms."""

# 5. Operating Systems (Different)
text_os = """Operating System Process Scheduling and Concurrency
Student Name: Vikram Singh
Roll Number: CS-2024-081

1. Role of Operating Systems
An operating system acts as an intermediary bridge between physical computer hardware and user software applications, providing process scheduling, virtual memory management, and file system abstractions.

2. Concurrency and Deadlocks
Deadlock occurs in concurrent computing when two or more execution threads are permanently halted because each process waits for a resource held by another process in a circular wait loop.

3. CPU Scheduling Algorithms
Common CPU scheduling policies include First-Come First-Served, Shortest Job First, Multi-Level Feedback Queues, and Preemptive Round Robin scheduling."""

# 6. Cloud Computing
text_cloud = """Cloud Computing Architectures and Containerization
Student Name: Neha Gupta
Roll Number: CS-2024-033

1. Cloud Computing Essentials
Cloud computing delivers on-demand computational capabilities including virtual servers, distributed object storage, managed database engines, and software services over high-speed networks.

2. Containerization and Microservices
Containers bundle application software along with its runtime system libraries into lightweight isolated user spaces. Tools like Docker and Kubernetes enable dynamic orchestration and self-healing deployments across heterogeneous clusters."""

# Write TXT files
with open(os.path.join(samples_dir, "Student_01_AI_Original.txt"), "w", encoding="utf-8") as f:
    f.write(text_ai_original)

with open(os.path.join(samples_dir, "Student_02_AI_Duplicate.txt"), "w", encoding="utf-8") as f:
    f.write(text_ai_duplicate)

with open(os.path.join(samples_dir, "Student_04_OS_Different.txt"), "w", encoding="utf-8") as f:
    f.write(text_os)

with open(os.path.join(samples_dir, "Student_05_Cloud_Different.txt"), "w", encoding="utf-8") as f:
    f.write(text_cloud)

# Write DOCX files
doc_para = Document()
for line in text_ai_paraphrased.split("\n"):
    if line.strip():
        doc_para.add_paragraph(line.strip())
doc_para.save(os.path.join(samples_dir, "Student_03_AI_Paraphrased.docx"))

doc_dbms = Document()
for line in text_dbms.split("\n"):
    if line.strip():
        doc_dbms.add_paragraph(line.strip())
doc_dbms.save(os.path.join(samples_dir, "Student_06_DBMS_Different.docx"))

print(f"Sample test documents created successfully in: {samples_dir}")
